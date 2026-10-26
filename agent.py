"""Agent 编排层：介于 NapCat 与 LLM 之间的"大脑 + 身体"。

职责：
1. 输入归一化：文本直达；图片→视觉描述；语音→ASR 转写；文件→下载并标注。
2. 工具循环：把归一化后的文本交给 LLM，LLM 可通过 tool_calling 调度
   web_search / generate_image / text_to_speech 等能力（由 services 真正执行）。
3. 记忆维护：把本轮对话写入 context_store，必要时触发摘要压缩。
4. 输出拼装：把 LLM 的回复与生成的图片/语音/文件，组装成 OneBot 消息段列表，
   交回 main 发送。

LLM 本身只说文本；图片/语音/文件都是 agent 调用 services 这个"手"产出的。
"""
import asyncio
import json
import logging
import re
import shutil
import time
from pathlib import Path

import services as svc
from ai_client import LLMClient

logger = logging.getLogger("QQAIBot.Agent")

# 媒体编号表：kind（内部类型）→ 对话里显示的中文标签
_MEDIA_LABEL = {"image": "图片", "video": "视频", "audio": "音频", "document": "文档"}
# 表里每个会话最多保留多少条媒体引用（超出按最旧淘汰，防止长期运行无限膨胀）
_MEDIA_MAX_ITEMS = 60

# ------------------------- 高精度解析库（可选） -------------------------
# 优先用专业库做高精度解析；若运行环境无法加载（如 alpha 版 Python 缺 lxml），
# 则回退到标准库 xml.etree 的 OOXML 解析，保证"能解析且不崩溃"。
try:
    from pypdf import PdfReader
    _HAVE_PYPDF = True
except Exception:  # noqa: BLE001
    _HAVE_PYPDF = False

try:
    import docx as _docx_mod
    _HAVE_DOCX = True
except Exception:  # noqa: BLE001
    _HAVE_DOCX = False

try:
    import openpyxl as _openpyxl_mod
    _HAVE_OPENPYXL = True
except Exception:  # noqa: BLE001
    _HAVE_OPENPYXL = False

try:
    from pptx import Presentation as _PptxPresentation
    _HAVE_PPTX = True
except Exception:  # noqa: BLE001
    _HAVE_PPTX = False

MAX_TOOL_ROUNDS = 6
# 各工具的必填参数。缺失或为空时**不执行**，改为回一条说明原因的 tool 消息，
# 让模型自己重发一次正确的参数——空参数执行的代价远大于空转一轮：
# web_search("") 会挨个引擎等到超时（最坏约 80 秒），generate_image("") 会真的出图并按图片发出。
# inspect_media 不在表内：它自带空值守卫（返回"需要提供媒体编号"），那句提示本身就是给模型看的指令。
_TOOL_REQUIRED = {
    "web_search": "query",
    "fetch_url": "url",
    "generate_image": "prompt",
    "text_to_speech": "text",
}


def _parse_tool_args(raw, name: str = "", finish_reason: str = "") -> tuple[dict, str]:
    """把 tool_call 的 arguments 解析成 dict，返回 (参数, 不可执行的原因)。

    第二项为空字符串表示可以执行；非空表示这次调用不该发出去，由调用方回一条
    tool 消息把原因原样告诉模型——模型据此重发正确参数，才是多轮工具循环该有的用法。
    只静默塞一个 {} 再照常执行，等于把"参数坏了"伪装成"工具正常返回空结果"，
    模型既不会重试，也无从知道该改什么。
    """
    if isinstance(raw, dict):
        # 少数网关/代理直接给对象而不是 JSON 字符串，能直接用就不必再解析
        args = raw
    elif raw is None or (isinstance(raw, str) and not raw.strip()):
        # 模型漏写参数：按空参处理，交给下面的必填校验判定
        args = {}
    elif not isinstance(raw, str):
        return {}, f"参数类型异常（{type(raw).__name__}，应为 JSON 字符串）"
    else:
        try:
            args = json.loads(raw)
        except ValueError:
            # 输出撞到 max_tokens 上限时，arguments 会停在半路，必然是坏 JSON
            why = ("参数不是合法 JSON（本次输出达到长度上限，被截断）"
                   if finish_reason == "length"
                   else "参数不是合法 JSON（可能是转义有误或输出不完整）")
            return {}, why

    if not isinstance(args, dict):
        return {}, f"参数必须是 JSON 对象（实际收到 {type(args).__name__}）"

    need = _TOOL_REQUIRED.get(name)
    if need:
        val = args.get(need)
        if val is None or (isinstance(val, str) and not val.strip()):
            return {}, f"缺少必需参数 {need}"
    return args, ""


def _reason(exc: BaseException) -> str:
    """把异常翻成一句能写进提示词的人话（不回显原始异常串）。

    原始 str(exc) 里可能带 URL、路径甚至请求细节，直接喂给模型容易被它复述给
    用户；这里只取"哪一类失败"，够模型判断"这块没拿到"就行。
    """
    if isinstance(exc, (asyncio.TimeoutError, TimeoutError)) or "timed out" in str(exc).lower():
        return "等待超时"
    if isinstance(exc, svc.ServiceUnavailable):
        # ServiceUnavailable 的文案本来就是人话（"未配置 api_key"等），去掉冒号后的细节
        return (str(exc).split(":")[0].strip() or "依赖服务不可用")[:40]
    if isinstance(exc, OSError):
        return "文件或网络读写失败"
    if isinstance(exc, (KeyError, IndexError, TypeError, ValueError)):
        return "消息内容格式异常"
    return "内部处理异常"


def _ensure_tool_call_ids(tool_calls: list, round_no: int) -> None:
    """给缺 id 的 tool_call 就地补一个本地 id。

    OpenAI 兼容协议要求：assistant 消息里每个 tool_call 都带 id，随后那条
    role="tool" 消息必须用同一个 id 回它。上游偶尔会漏这个字段（网关改写、自建
    兼容层、非标准实现），不补的话下一轮请求会被判"消息结构不合法"直接 400，
    整轮回复连同一轮工具结果一起丢掉，日志里只能看到一句含糊的接口报错。

    补出来的 id 只需在本轮上下文里自洽——它不发给真正的工具执行方，上游也
    不会去校验格式之外的任何东西。已有 id 的一律不动。
    """
    for i, tc in enumerate(tool_calls):
        if isinstance(tc, dict) and not tc.get("id"):
            tc["id"] = f"call_local_{round_no}_{i}"


# 音频类文件：交给 ASR / 音频模型（录音文件识别，异步提交+轮询；或方舟 Chat 音频理解）
_AUDIO_EXTS = {".mp3", ".wav", ".m4a", ".amr", ".ogg", ".flac", ".aac", ".wma",
               ".opus", ".silk"}
# 视频类文件：交给豆包原生视频理解（容器型扩展名；音频专用扩展名留在上面）
_VIDEO_EXTS = {".mp4", ".mov", ".avi", ".mkv", ".webm", ".flv", ".wmv", ".m4v",
               ".3gp", ".mpg", ".mpeg", ".ts"}
# 视频直链扩展名（文本里的链接若以此结尾，视为可直接喂给视频模型，而非抓网页）
_VIDEO_URL_EXTS = {".mp4", ".mov", ".avi", ".mkv", ".webm", ".flv", ".wmv", ".m4v",
                   ".3gp", ".mpg", ".mpeg"}
# 图片类文件：交给视觉理解
_IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp", ".tiff",
               ".heic", ".heif"}
# 文档类（交给方舟 Chat 的 file 内容类型做多模态文档理解，走 document_model；失败回退本地解析）
_DOC_MODEL_EXTS = {".pdf", ".doc", ".docx", ".xls", ".xlsx", ".ppt", ".pptx"}
# 文本类文件：直接本地读取内容（无需模型）
_TEXT_LIKE = {".txt", ".md", ".py", ".json", ".csv", ".log", ".yaml", ".yml",
              ".xml", ".html", ".js", ".ts", ".java", ".c", ".cpp", ".go", ".rs",
              ".toml", ".ini", ".cfg"}
# 链接里直接以这些扩展名结尾 → 视为对应媒体直链（走模型而非抓网页）
_IMAGE_URL_EXTS = _IMAGE_EXTS
# 音频直链扩展名（.silk 是 QQ 私有格式，不会出现在公网 URL 里，故不收录）
_AUDIO_URL_EXTS = {".mp3", ".wav", ".m4a", ".aac", ".ogg", ".flac",
                   ".opus", ".amr", ".wma"}
# Office 类：用 zipfile 简易抽取文本（_DOC_MODEL_EXTS 的本地解析兜底用）
_OFFICE_EXTS = {".docx", ".xlsx", ".pptx"}


def _cleanup_downloaded(path: str) -> None:
    """清理 yt-dlp 下载的临时视频（文件 + 它所在的 qqbot-media-* 临时目录）。"""
    try:
        p = Path(path)
        parent = p.parent
        if p.exists():
            p.unlink()
        if parent.name.startswith(("qqbot-media-", "qqbot-doc-")) and parent.is_dir():
            shutil.rmtree(parent, ignore_errors=True)
    except Exception as exc:  # noqa: BLE001
        logger.warning("清理临时视频失败 %s: %s", path, exc)


# ------------------------------------------------------------------ PDF
def _read_pdf(path: str, max_chars: int):
    """高精度：pypdf 逐页抽取；缺失/失败回退到字节流正则兜底。"""
    if _HAVE_PYPDF:
        try:
            with PdfReader(path) as reader:
                parts = [ (page.extract_text() or "") for page in reader.pages ]
            text = "\n".join(parts).strip()
            if text:
                return text[:max_chars]
            # 可能是扫描件/图片型 PDF（无文本层）
            return "（PDF 未提取到文本：可能是扫描件/图片型 PDF，需要 OCR 才能识别）"
        except Exception as exc:  # noqa: BLE001
            logger.warning("pypdf 解析失败 %s: %s，回退简易提取", path, exc)
    return _read_pdf_simple(path, max_chars)


def _read_pdf_simple(path: str, max_chars: int):
    import re as _re
    try:
        raw = Path(path).read_bytes()
        strings = _re.findall(rb"\(((?:[^()\\]|\\.)*)\)", raw)
        pieces = []
        for s in strings:
            t = (s.replace(b"\\(", b"(").replace(b"\\)", b")")
                   .replace(b"\\\\", b" ")).decode("latin-1", "ignore")
            if _re.search(r"[\u4e00-\u9fffA-Za-z]", t):
                pieces.append(t)
        text = _re.sub(r"\s+", " ", " ".join(pieces)).strip()
        return ("（PDF 为简易提取，可能不完整）\n" + text)[:max_chars] if text else None
    except Exception:  # noqa: BLE001
        return None


# ------------------------------------------------------------------ Office
def _read_docx(path: str, max_chars: int):
    if _HAVE_DOCX:
        try:
            d = _docx_mod.Document(path)
            paras = [p.text for p in d.paragraphs if p.text and p.text.strip()]
            for tbl in d.tables:
                for row in tbl.rows:
                    cells = [c.text for c in row.cells if c.text and c.text.strip()]
                    if cells:
                        paras.append(" | ".join(cells))
            text = "\n".join(paras).strip()
            return text[:max_chars] if text else None
        except Exception as exc:  # noqa: BLE001
            logger.warning("python-docx 解析失败 %s: %s，回退标准库解析", path, exc)
    return _read_office_xml(path, max_chars, "docx")


def _read_xlsx(path: str, max_chars: int):
    if _HAVE_OPENPYXL:
        try:
            wb = _openpyxl_mod.load_workbook(path, read_only=True, data_only=True)
            try:
                lines = []
                for ws in wb.worksheets:
                    for row in ws.iter_rows(values_only=True):
                        cells = [str(c) for c in row if c is not None and str(c).strip()]
                        if cells:
                            lines.append(" | ".join(cells))
                text = "\n".join(lines).strip()
                return text[:max_chars] if text else None
            finally:
                wb.close()
        except Exception as exc:  # noqa: BLE001
            logger.warning("openpyxl 解析失败 %s: %s，回退标准库解析", path, exc)
    return _read_office_xml(path, max_chars, "xlsx")


def _read_pptx(path: str, max_chars: int):
    if _HAVE_PPTX:
        try:
            prs = _PptxPresentation(path)
            lines = []
            for slide in prs.slides:
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for para in shape.text_frame.paragraphs:
                            t = "".join(r.text for r in para.runs).strip()
                            if t:
                                lines.append(t)
                    if shape.has_table:
                        for row in shape.table.rows:
                            cells = [c.text for c in row.cells if c.text and c.text.strip()]
                            if cells:
                                lines.append(" | ".join(cells))
            text = "\n".join(lines).strip()
            return text[:max_chars] if text else None
        except Exception as exc:  # noqa: BLE001
            logger.warning("python-pptx 解析失败 %s: %s，回退标准库解析", path, exc)
    return _read_office_xml(path, max_chars, "pptx")


def _read_office_xml(path: str, max_chars: int, kind: str):
    """标准库兜底：用 zipfile + xml.etree 解析 OOXML（比旧版正则精确得多）。

    覆盖 docx/xlsx/pptx 的通用命名空间，逐元素抽取文本；当专业库缺失
    （如 alpha 版 Python 无法加载 lxml）时作为主力解析器使用。
    """
    import zipfile
    import xml.etree.ElementTree as ET
    try:
        with zipfile.ZipFile(path) as z:
            if kind == "docx":
                W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
                root = ET.fromstring(z.read("word/document.xml"))
                texts = []
                for p in root.iter(W + "p"):
                    txt = "".join(t.text or "" for t in p.iter(W + "t"))
                    if txt.strip():
                        texts.append(txt)
                out = "\n".join(texts)
            elif kind == "pptx":
                A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
                texts = []
                for n in z.namelist():
                    if ((n.startswith("ppt/slides/slide") or
                         n.startswith("ppt/notesSlides/notesSlide")) and n.endswith(".xml")):
                        root = ET.fromstring(z.read(n))
                        for t in root.iter(A + "t"):
                            if t.text and t.text.strip():
                                texts.append(t.text)
                out = "\n".join(texts)
            else:  # xlsx
                M = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
                shared = {}
                if "xl/sharedStrings.xml" in z.namelist():
                    sroot = ET.fromstring(z.read("xl/sharedStrings.xml"))
                    for i, si in enumerate(sroot.iter(M + "si")):
                        shared[i] = "".join(t.text or "" for t in si.iter(M + "t"))
                lines = []
                for n in z.namelist():
                    if n.startswith("xl/worksheets/sheet") and n.endswith(".xml"):
                        root = ET.fromstring(z.read(n))
                        for c in root.iter(M + "c"):
                            v = c.find(M + "v")
                            if v is None or v.text is None:
                                continue
                            if c.get("t") == "s":
                                lines.append(shared.get(int(v.text), ""))
                            else:
                                lines.append(v.text)
                out = "\n".join(lines)
        out = out.strip()
        return out[:max_chars] if out else None
    except Exception as exc:  # noqa: BLE001
        logger.warning("标准库解析 Office(%s) 失败 %s: %s", kind, path, exc)
        return None


def _read_file_text(path: str, max_chars: int = 6000):
    """高精度从文件中抽取可读文本；不支持的类型返回 None。

    文本类直接读；PDF 用 pypdf（扫描件提示用 OCR）；docx/xlsx/pptx 优先用
    python-docx/openpyxl/python-pptx，缺失时回退标准库 OOXML 解析。
    """
    p = Path(path)
    if not p.exists():
        return None
    ext = p.suffix.lower()
    try:
        if ext in _TEXT_LIKE:
            return p.read_text(encoding="utf-8", errors="ignore")[:max_chars]
        if ext == ".pdf":
            return _read_pdf(path, max_chars)
        if ext in _OFFICE_EXTS:
            if ext == ".docx":
                return _read_docx(path, max_chars)
            if ext == ".xlsx":
                return _read_xlsx(path, max_chars)
            return _read_pptx(path, max_chars)
    except Exception as exc:  # noqa: BLE001
        logger.warning("读取文件文本失败 %s: %s", path, exc)
    return None


def _tool_specs(openai_ready: bool) -> list:
    specs = [{
        "type": "function",
        "function": {
            "name": "web_search",
            "description": "联网搜索实时/最新信息（新闻、天气、股价、资料等）。当需要最新或不确定信息时调用。",
            "parameters": {
                "type": "object",
                "properties": {"query": {"type": "string", "description": "搜索关键词"}},
                "required": ["query"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "fetch_url",
            "description": ("打开并读取一个网页/链接的内容（用户发来链接、或要你访问某个网址看内容时调用）。"
                            "返回页面纯文本；若页面内嵌视频/音频/图片，会自动用多模态模型分析并返回描述。"),
            "parameters": {
                "type": "object",
                "properties": {"url": {"type": "string", "description": "要打开的网址，如 https://example.com"}},
                "required": ["url"],
            },
        },
    }]
    if openai_ready:
        specs += [
            {
                "type": "function",
                "function": {
                    "name": "inspect_media",
                    "description": (
                        "重新查看本会话里某个已发过的图片/视频/音频/文档，针对细节提问。"
                        "当历史消息里出现 [图片 m1] / [视频 m2] 这类标记、而现有描述不足以回答时调用。"
                        "必须给出具体问题，例如「画面右下角的文字是什么」「第 3 秒出现了谁」；"
                        "禁止只写「再看看」「再仔细点」这类没有信息量的请求。"
                        "同一个媒体最多追问两次，得到答案后就停止。"),
                    "parameters": {
                        "type": "object",
                        "properties": {
                            "handle": {"type": "string",
                                       "description": "媒体编号，如 m1（见历史里的 [图片 m1] 标记）"},
                            "question": {"type": "string",
                                         "description": "要确认的具体细节，必须是可回答的具体问题"},
                        },
                        "required": ["handle", "question"],
                    },
                },
            },
            {
                "type": "function",
                "function": {
                    "name": "generate_image",
                    "description": "根据文字描述生成一张图片作为回复（用户要你画图、发图、做表情包时调用）。",
                    "parameters": {
                        "type": "object",
                        "properties": {"prompt": {"type": "string", "description": "画面描述，尽量具体"}},
                        "required": ["prompt"],
                    },
                },
            },
            {
                "type": "function",
                "function": {
                    "name": "text_to_speech",
                    "description": "把指定文本合成为语音并作为语音消息回复（用户要你用语音/发语音时调用）。",
                    "parameters": {
                        "type": "object",
                        "properties": {"text": {"type": "string", "description": "要念出的文本"}},
                        "required": ["text"],
                    },
                },
            },
        ]
    return specs


class Agent:
    def __init__(self, llm: LLMClient, services: "svc.Services", max_reply_length: int = 1500,
                 recover_timeout: float = 60):
        self.llm = llm
        self.services = services
        self.max_reply_length = max_reply_length
        # 补救回复的单次等待上限（秒）。刻意小于 event_handle：失败路径不该再吃满总预算
        self.recover_timeout = float(recover_timeout) if recover_timeout else 60.0
        if self.recover_timeout <= 0:
            self.recover_timeout = 60.0

    # ------------------------- 对外入口 -------------------------
    async def handle(self, session_id: str, user_input: dict,
                     context_store: dict, mem_cfg: dict, call_api) -> list:
        """处理一轮用户输入，返回要发送的 OneBot 消息段列表。

        user_input 结构：
          {"text": str, "images": [本地路径], "voices": [本地路径],
           "videos": [本地路径], "files": [(name, path), ...], "sender_name": str}
          —— files 为上传的文件（群文件/离线文件），按扩展名自动路由到
             视频理解 / 视觉理解 / 语音识别 / 文本抽取。
        """
        # 1) 输入归一化
        # 1) 会话上下文：先取到这一格，归一化时要把媒体引用登记进去（m1/m2…）
        entry = context_store.setdefault(session_id, {"summary": "", "recent": []})

        # 逐份媒体的失败已在 _normalize 内部逐项降级（变成 prompt 里一行"识别失败"）。
        # 这里兜的是更外层的事故（下载目录不可写、登记表结构坏了等）：同样不给用户
        # 甩固定文案，而是降级成"原文 + 哪些媒体没看成"，让模型照常作答。
        try:
            user_text = await self._normalize(user_input, call_api, entry)
        except Exception as exc:  # noqa: BLE001
            logger.exception("输入归一化失败，降级为纯文本: %s", exc)
            user_text = self._fallback_user_text(user_input, exc)

        # 2) 组装消息
        #    时间戳只加在本轮 user 消息上：它位于前缀最末尾，不会改变历史前缀，
        #    因此不影响 prompt 缓存命中；同时最后一条时间戳天然代表"现在"。
        stamp = f"[{time.strftime('%m-%d %H:%M')}]"
        user_msg = f"{stamp}\n{user_text}" if user_text else stamp
        messages = self._build_messages(entry, user_msg)

        # 3) 工具循环
        outputs = {"images": [], "voices": [], "files": []}
        tools = _tool_specs(self.services.openai_ready)
        try:
            reply_text = await self._tool_loop(messages, tools, outputs, entry, call_api)
        except Exception as exc:  # noqa: BLE001
            # LLM 自己的超时/限流/鉴权已在 _tool_loop 内转成固定回复，走不到这里；
            # 能到这儿的都是"意外"（工具层之外的 bug、IO 故障等）。先让模型带着
            # 已有上下文把话说完，比直接甩一句"处理出错了"有用得多。
            logger.exception("工具循环失败，尝试补救回复: %s", exc)
            reply_text = await self._recover_reply(entry, user_msg, exc)
            if reply_text is None:
                # 补救也没成：原样抛出，交给 main 的固定文案兜底（日志仍在 worker 里打）
                raise

        if not reply_text:
            # 上游偶发空响应：也要给一句能看懂的话。空 segments 会被 main 当成
            # "没什么要说的"直接丢弃，用户等半天一个字都收不到。
            reply_text = "（模型这次返回了空内容，可以再问一次。）"

        # 4) 写回记忆（存带时间戳的版本，历史里就能看出每句话是什么时候说的）
        entry.setdefault("recent", []).append({"role": "user", "content": user_msg})
        entry["recent"].append({"role": "assistant", "content": reply_text})
        mode = mem_cfg.get("mode", "accumulate")
        if mode == "sliding":
            limit = int(mem_cfg.get("max_history", 12)) * 2
            if len(entry["recent"]) > limit:
                del entry["recent"][: len(entry["recent"]) - limit]
        else:
            await self.llm.maybe_summarize(session_id, entry, mem_cfg, messages + [
                {"role": "user", "content": user_msg},
                {"role": "assistant", "content": reply_text},
            ])

        # 5) 拼装输出段
        return self._build_segments(reply_text, outputs)

    # ------------------------- 失败补救 -------------------------
    def _fallback_user_text(self, user_input: dict, exc: BaseException) -> str:
        """归一化整体失败时的降级文本：原文照留，只标注哪些媒体这次没看成。"""
        text = (user_input.get("text") or "").strip()
        marks = []
        for key, label in (("images", "图片"), ("voices", "语音"),
                           ("videos", "视频"), ("files", "文件")):
            n = len(user_input.get(key) or [])
            if n:
                marks.append(f"{label}×{n}")
        note = ""
        if marks:
            note = (f"[本轮没能处理用户发来的{'、'.join(marks)}"
                    f"（{_reason(exc)}），只能按文字理解]")
        return "\n".join(p for p in (text, note) if p)

    async def _recover_reply(self, entry: dict, user_msg: str,
                             exc: BaseException) -> str | None:
        """中途失败后的一次补救：带着已有上下文请模型把话说完。

        只用于非 LLM 类失败（LLM 超时/限流/鉴权由 _tool_loop 自己消化）。
        返回 None 表示补救也没成，调用方应原样抛出、让 main 的固定文案兜底。
        """
        reason = _reason(exc)
        messages = self._build_messages(entry, user_msg)
        messages.append({"role": "system", "content": (
            f"【系统提示】本轮处理中途失败：{reason}。"
            "请基于上面的对话和已经成功获取的信息，回答用户最后一条消息；"
            "没拿到的部分（某个附件、某个链接、某个工具结果）直接说明没成功，"
            "不要猜测或编造其中的内容。")})
        try:
            # tools=None：补救只求一句能发出去的话，不再触发工具，避免二次踩坑。
            # wait_for 只约束"我们还等不等"，底层线程要等自己的 socket 超时才退出，
            # 但这是罕见的失败路径，可接受。
            resp = await asyncio.wait_for(
                self.llm.complete(messages, None), timeout=self.recover_timeout)
        except Exception as exc2:  # noqa: BLE001
            logger.warning("补救回复失败（不再重试）: %s", exc2)
            return None
        choices = resp.get("choices") or []
        content = (((choices[0] if choices else {}) or {}).get("message") or {}).get("content") or ""
        if not content.strip():
            logger.warning("补救回复为空（choices=%d）: %s", len(choices), str(resp)[:200])
            return None
        logger.info("补救回复成功（原始失败：%s）", reason)
        return content.strip()

    # ------------------------- 媒体引用登记（m1 / m2 …） -------------------------
    def _next_handle(self, entry: dict) -> str:
        """会话内自增编号：只增不减、永不复用。

        复用编号会造成"m1 上一轮是猫、这一轮是狗"，历史文本里同一个 m1 指向不同资源，
        模型必然串味。所以宁可让数字变大，也不回收。
        """
        raw = entry.get("media_seq")
        if raw is None:
            # 兼容旧记忆文件（没有 media_seq）：从已有编号推最大值，避免撞号
            nums = [int(k[1:]) for k in (entry.get("media") or {})
                    if isinstance(k, str) and k[1:].isdigit()]
            raw = max(nums) if nums else 0
        n = int(raw) + 1
        entry["media_seq"] = n
        return f"m{n}"

    def _register_media(self, entry: dict, kind: str, ref: dict | None, *,
                        msg_id=None, seg=None,
                        audio_ref: dict | None = None) -> str | None:
        """把一次上传的引用登记进会话表，返回对话里使用的短编号（如 m1）。

        表里只存**指针**：file_id / url / 音轨 file_id / 可重下的消息定位。
        **不存描述**——描述已经在 recent 的文本里，存两份会出现"谁为准"的问题。
        """
        if entry is None or not ref:
            return None
        fid = (ref or {}).get("file_id") or ""
        url = (ref or {}).get("url") or (ref or {}).get("file_url") or ""
        if url.startswith("data:"):
            url = ""                     # 内联 base64 是几十 MB 字符串，不值得存
        if not fid and not url and not msg_id:
            return None                  # 既不能复用也不能重下，登记了没意义

        h = self._next_handle(entry)
        now = int(time.time())
        days = int(getattr(self.services, "media_expire_days", 0) or 0) or 7
        rec = {"kind": kind, "created": now, "expire_at": now + days * 86400}
        if fid:
            rec["file_id"] = fid
        if url:
            rec["url"] = url
        if audio_ref and audio_ref.get("file_id"):
            rec["audio_file_id"] = audio_ref["file_id"]
        if msg_id:
            rec["msg_id"] = str(msg_id)
            if seg is not None:
                rec["seg"] = int(seg)

        entry.setdefault("media", {})[h] = rec
        self._trim_media(entry)
        logger.info("登记媒体 %s：%s（file_id=%s msg_id=%s）",
                    h, kind, fid or "-", rec.get("msg_id") or "-")
        return h

    def _trim_media(self, entry: dict) -> None:
        """先剔除"已过期且没有重下线索"的条目，再按上限淘汰最旧的。"""
        table = entry.get("media") or {}
        now = time.time()
        for k in [k for k, v in table.items()
                  if not v.get("msg_id") and float(v.get("expire_at") or 0) < now]:
            table.pop(k, None)
        if len(table) > _MEDIA_MAX_ITEMS:
            for k in list(table)[: len(table) - _MEDIA_MAX_ITEMS]:
                table.pop(k, None)

    @staticmethod
    def _ref_of(rec: dict) -> dict:
        if rec.get("file_id"):
            return {"file_id": rec["file_id"]}
        if rec.get("url"):
            return {"url": rec["url"]}
        return {}

    @staticmethod
    def _audio_ref_of(rec: dict) -> dict | None:
        fid = rec.get("audio_file_id")
        return {"file_id": fid} if fid else None

    @staticmethod
    def _ask_prompt(question: str) -> str:
        return ("请针对下面的具体问题仔细观察这个媒体并回答；"
                "如果其中确实看不到相关信息，请明确说「看不到」，不要凭常理猜测。\n"
                f"问题：{question}")

    async def _ensure_media_ref(self, rec: dict, handle: str, call_api):
        """三级解析：远端未过期 → 远端校验 → 回 QQ 重下。

        返回 (ref, audio_ref, local_path, err)：
          ref 非空   → 可直接用已存在的 file_id/url 追问（不重新上传）
          local 非空 → 已从 QQ 重新下载到本地，调用方需重新分析并回写新 file_id
          err 非空   → 彻底不可用
        """
        fid = rec.get("file_id")

        # ① 本地时间判断：离过期还早，直接复用（覆盖绝大多数追问，零额外请求）
        if fid and time.time() < float(rec.get("expire_at") or 0) - 3600:
            return {"file_id": fid}, self._audio_ref_of(rec), None, None

        # ② 远端校验：可能被 cleanup 脚本删了，或本地时间戳不准
        if fid and await asyncio.to_thread(self.services.file_exists, fid):
            # 只往前推 1 天，不能重新按 30 天算——远端过期时间不会因为查询而延长，
            # 否则会在文件真正过期后仍被当成"有效"，追问时才突然 404。
            rec["expire_at"] = int(time.time()) + 86400
            return {"file_id": fid}, self._audio_ref_of(rec), None, None

        # ③ 回 QQ 重新下载（NapCat 会重新签发 CDN 链接）
        if rec.get("msg_id") and call_api is not None:
            local = await self._refetch_from_qq(rec, handle, call_api)
            if local:
                return None, None, local, None

        # ④ 只剩一条直链可用就用它；否则彻底失败
        if rec.get("url"):
            return {"url": rec["url"]}, None, None, None
        return None, None, None, "远端文件已过期，QQ 侧也未能重新取到"

    async def _refetch_from_qq(self, rec: dict, handle: str, call_api) -> str | None:
        """用 msg_id + 同类段序号，从 QQ 重新取回原始媒体到本地。

        只覆盖图片/视频/语音：文档走的是另一套（file 段 + busid），
        且体积小、用户重发成本低，过期后直接如实告知即可。
        """
        import media as media_mod
        kind = (rec.get("kind") or "image").lower()
        seg_type = {"image": "image", "video": "video", "audio": "record"}.get(kind)
        if seg_type is None:
            return None
        want = int(rec.get("seg") or 0)
        try:
            resp = await call_api("get_msg", {"message_id": int(rec["msg_id"])})
        except Exception as exc:  # noqa: BLE001
            logger.warning("重下 %s 失败（get_msg）: %s", handle, exc)
            return None
        segs = ((resp or {}).get("data") or {}).get("message") or []
        target, k = None, 0
        for s in segs:
            if not isinstance(s, dict) or s.get("type") != seg_type:
                continue
            if k == want:
                target = s
                break
            k += 1
        if target is None:
            logger.warning("重下 %s 失败：消息 %s 里没有第 %d 个 %s 段",
                           handle, rec.get("msg_id"), want, seg_type)
            return None
        data = target.get("data") or {}
        try:
            if kind == "image":
                local = await media_mod.fetch_image(data, call_api)
            elif kind == "video":
                local = await media_mod.fetch_video(data, call_api)
            else:
                local = await media_mod.fetch_voice(data, call_api)
        except Exception as exc:  # noqa: BLE001
            logger.warning("重下 %s 下载失败: %s", handle, exc)
            return None
        if local:
            logger.info("媒体 %s 已从 QQ 重新下载：%s", handle, Path(local).name)
        return local

    async def _ask_local(self, local: str, kind: str, question: str):
        """用重新下载到的本地文件分析，并顺便把新引用带出来（用于回写）。"""
        sink: dict = {}
        prompt = self._ask_prompt(question)
        if kind == "image":
            ans = await self.services.describe_image(local, prompt=prompt, sink=sink)
        elif kind == "video":
            ans = await self.services.understand_video(local, prompt=prompt, sink=sink)
        else:
            ans = await self.services.understand_audio(local, prompt=prompt, sink=sink)
        return ans, sink

    async def _inspect_media(self, entry: dict, handle: str, question: str,
                             call_api) -> str:
        """工具入口：按编号重新查看某个媒体并回答具体问题。"""
        raw = str(handle or "").strip()
        question = (question or "").strip()
        if not raw:
            return "（需要提供媒体编号，例如 m1。编号见历史消息里的 [图片 m1] 标记。）"
        if not question:
            return "（需要具体问题，例如「画面右下角的文字是什么」。不允许只说「再看看」。）"

        table = entry.get("media") or {}
        key = raw.lower()
        rec = table.get(key)
        if rec is None:
            # 容错：模型可能写成 "图片 m1" / "[m1]" / "m1:" 等
            blob = re.sub(r"[\[\]（）()：:\s]", "", key)
            for k, v in table.items():
                if k and k in blob:
                    key, rec = k, v
                    break
        if rec is None:
            avail = "、".join(list(table)[-6:]) or "无"
            return f"（会话里没有 {raw} 这个媒体。当前可用编号：{avail}）"

        kind = (rec.get("kind") or "image").lower()
        ref, audio_ref, local, err = await self._ensure_media_ref(rec, key, call_api)
        try:
            if local:
                ans, sink = await self._ask_local(local, kind, question)
                new_ref = sink.get("ref") or {}
                if new_ref.get("file_id"):
                    days = int(getattr(self.services, "media_expire_days", 0) or 0) or 7
                    rec["file_id"] = new_ref["file_id"]
                    rec["created"] = int(time.time())
                    rec["expire_at"] = rec["created"] + days * 86400
                    logger.info("媒体 %s 已重新上传，file_id=%s", key, rec["file_id"])
                new_audio = sink.get("audio_ref") or {}
                if new_audio.get("file_id"):
                    rec["audio_file_id"] = new_audio["file_id"]
                return ans
            if err:
                return (f"（{key} 已不可用：{err}。如果只需确认之前描述过的内容，"
                        f"可以基于上下文回答；否则请让用户重发一次。）")
            if not ref:
                return (f"（{key} 没有可复用的远端引用，也没有本地文件可重传，"
                        f"请让用户重发一次。）")
            return await self.services.ask_media(ref, kind, question, audio_ref)
        except svc.ServiceUnavailable as exc:
            return f"（追问 {key} 失败：{exc}）"

    # ------------------------- 媒体链接路由 -------------------------
    async def _describe_media_url(self, url: str) -> str | None:
        """若 URL 是媒体直链，交给对应多模态模型"看/听"，返回可直接入 prompt 的文本。

        覆盖四类（按扩展名判定，忽略 query/fragment）：
          视频 → video_model（understand_video）    图片 → vision_model（describe_image）
          文档 → document_model（understand_document） 音频 → audio_model（understand_audio）

        返回 None 表示"不是媒体直链"或"模型不可用/调用失败"，由调用方回退普通抓取，
        这样媒体链接和普通网页共用同一套兜底逻辑，不会因为一次模型失败就丢掉链接。
        """
        lower = url.lower().split("?")[0].split("#")[0]
        try:
            # 平台视频页面（B站/YouTube/抖音…）：本身不是媒体直链，
            # 但可以用 yt-dlp 取出视频流真正"看画面"；失败则回退网页抓取
            if svc._is_video_page(url) and self.services.video_ready:
                desc = await self._describe_platform_video(url)
                if desc is not None:
                    return desc

            if any(lower.endswith(ext) for ext in _VIDEO_URL_EXTS):
                if not self.services.video_ready:
                    logger.warning("视频链接但视频模型未就绪: %s", url)
                    return f"[用户发送了视频链接 {url}，但当前未配置视频理解服务，无法分析画面]"
                desc = await self.services.understand_video(url, is_url=True)
                return f"[视频链接内容 {url}]\n{desc}"

            if any(lower.endswith(ext) for ext in _IMAGE_URL_EXTS):
                if not self.services.openai_ready:
                    logger.warning("图片链接但视觉模型未就绪: %s", url)
                    return f"[用户发送了图片链接 {url}，但当前未配置视觉服务，无法识别]"
                desc = await self.services.describe_image(url, is_url=True)
                return f"[图片链接内容 {url}]\n{desc}"

            if any(lower.endswith(ext) for ext in _DOC_MODEL_EXTS):
                if not self.services.document_model:
                    return None  # 未配置文档模型 → 回退抓网页文本
                desc = await self._read_document_url(url)
                return f"[文档链接内容 {url}]\n{desc}"

            if any(lower.endswith(ext) for ext in _AUDIO_URL_EXTS):
                if not self.services.audio_model:
                    logger.warning("音频链接但音频模型未就绪: %s", url)
                    return f"[用户发送了音频链接 {url}，但当前未配置音频理解服务，无法识别]"
                desc = await self.services.understand_audio(url, is_url=True)
                return f"[音频链接内容 {url}]\n{desc}"

            # 扩展名一个都没命中：用 Content-Type 探测一次，
            # 覆盖 /download?id=123 这类"路径里没有扩展名"的媒体/文档直链
            cfg = self.services.web_config or {}
            if cfg.get("probe_content_type", True):
                ctype = await asyncio.to_thread(self.services.probe_content_type, url)
                if ctype:
                    if svc._is_doc_mime(ctype):
                        if self.services.document_model:
                            desc = await self._read_document_url(url)
                            return f"[文档链接内容 {url}]\n{desc}"
                    elif ctype.startswith("video/") and self.services.video_ready:
                        desc = await self.services.understand_video(url, is_url=True)
                        return f"[视频链接内容 {url}]\n{desc}"
                    elif ctype.startswith("audio/") and self.services.audio_model:
                        desc = await self.services.understand_audio(url, is_url=True)
                        return f"[音频链接内容 {url}]\n{desc}"
                    elif ctype.startswith("image/") and self.services.openai_ready:
                        desc = await self.services.describe_image(url, is_url=True)
                        return f"[图片链接内容 {url}]\n{desc}"
        except svc.ServiceUnavailable as exc:
            logger.warning("媒体链接理解失败，回退网页抓取 %s: %s", url, exc)
            return None
        except Exception as exc:  # noqa: BLE001
            logger.warning("媒体链接处理异常，回退网页抓取 %s: %s", url, exc)
            return None
        return None  # 非媒体直链

    async def _describe_platform_video(self, url: str) -> str | None:
        """视频平台页面（B站/YouTube/抖音…）→ yt-dlp 下载 → 视频模型看画面。

        返回 None 表示"没拿到"（下载失败/超时长/超体积/开关关闭），由调用方回退到
        fetch_url 抓网页拿标题简介——下载不了也不至于什么都给不出。
        """
        cfg = self.services.web_config or {}
        if cfg.get("page_media", True) is False:
            return None
        try:
            path, meta = await asyncio.to_thread(
                self.services.download_media, url,
                int(cfg.get("media_download_max_seconds", 180)),
                int(cfg.get("media_download_max_mb", 40)))
        except svc.ServiceUnavailable as exc:
            logger.warning("平台视频未下载，回退网页抓取 %s: %s", url, exc)
            return None
        try:
            desc = await self.services.understand_video(path)
        except svc.ServiceUnavailable as exc:
            logger.warning("平台视频理解失败 %s: %s", url, exc)
            return None
        finally:
            _cleanup_downloaded(path)
        title = (meta.get("title") or "").strip()
        dur = float(meta.get("duration") or 0)
        label = f"《{title}》" if title else ""
        dur_s = f" 时长{int(dur)}秒" if dur else ""
        return f"[视频页面{label}{dur_s} {url}]\n{desc}"

    async def _read_document_url(self, url: str) -> str:
        """读一个文档直链：先交给方舟按 file_url 拉取，失败再本机下载后走 Files API。

        为什么需要两步：不少站点的文档直链有防盗链/重定向/奇怪 MIME，方舟服务端
        直接拉会 400；而本机带上 UA 下载通常没问题，落地后再上传就稳了。
        """
        try:
            return await self.services.understand_document(url, is_url=True)
        except svc.ServiceUnavailable as exc:
            logger.warning("文档直链拉取失败，改本机下载 %s: %s", url, exc)
        local = await asyncio.to_thread(self.services.download_file, url, 30)
        try:
            return await self.services.understand_document(local)
        finally:
            _cleanup_downloaded(local)

    async def _describe_page_media(self, url: str) -> str | None:
        """打开网页，把它**内嵌的媒体**交给模型"看/听"，返回可入 prompt 的文本。

        与 _describe_media_url（处理媒体直链）互补：这里处理的是"网页里装着媒体"
        的情况——新闻页里的视频、播客页的音频、文章主图、B站/YouTube 播放器。

        取用策略（控成本）：最多 1 个视频（平台播放器优先，其次直链）+ 1 个音频
        + media_max_images 张图片；都取不到返回 None，调用方只保留网页文本。
        单个媒体失败只跳过它自己（append 一条说明），不影响其它媒体。
        """
        cfg = self.services.web_config or {}
        if cfg.get("page_media", True) is False:
            return None
        try:
            media = await self.services.fetch_page_media(url, web_cfg=cfg, max_each=3)
        except Exception as exc:  # noqa: BLE001
            logger.warning("提取页面媒体失败 %s: %s", url, exc)
            return None

        videos = media.get("videos") or []
        audios = media.get("audios") or []
        images = media.get("images") or []
        embeds = media.get("embeds") or []
        docs = media.get("docs") or []
        if not (videos or audios or images or embeds or docs):
            return None

        notes: list[str] = []
        left_videos = int(cfg.get("media_max_videos", 1))

        # 1) 平台播放器（B站/YouTube/抖音…）→ yt-dlp 下载后真正"看画面"
        if embeds and left_videos > 0 and self.services.video_ready:
            target = embeds[0]
            try:
                path, meta = await asyncio.to_thread(
                    self.services.download_media, target,
                    int(cfg.get("media_download_max_seconds", 180)),
                    int(cfg.get("media_download_max_mb", 40)))
            except svc.ServiceUnavailable as exc:
                logger.warning("页面内嵌视频未下载 %s: %s", target, exc)
                notes.append(f"[页面内嵌视频 {target} 未能分析：{exc}]")
            else:
                try:
                    desc = await self.services.understand_video(path)
                    title = (meta.get("title") or "").strip()
                    label = f"《{title}》" if title else ""
                    notes.append(f"[页面内视频{label} {target}]\n{desc}")
                    left_videos -= 1
                except svc.ServiceUnavailable as exc:
                    logger.warning("页面内嵌视频理解失败 %s: %s", target, exc)
                    notes.append(f"[页面内嵌视频 {target} 未能分析：{exc}]")
                finally:
                    _cleanup_downloaded(path)

        # 2) 页面里直接给出的视频文件地址
        if videos and left_videos > 0 and self.services.video_ready:
            try:
                desc = await self.services.understand_video(videos[0], is_url=True)
                notes.append(f"[页面内视频 {videos[0]}]\n{desc}")
            except svc.ServiceUnavailable as exc:
                logger.warning("页面视频直链理解失败 %s: %s", videos[0], exc)

        # 3) 页面里嵌的音频
        if audios and self.services.audio_model:
            try:
                desc = await self.services.understand_audio(audios[0], is_url=True)
                notes.append(f"[页面内音频 {audios[0]}]\n{desc}")
            except svc.ServiceUnavailable as exc:
                logger.warning("页面音频理解失败 %s: %s", audios[0], exc)

        # 4) 页面主图（og:image / 大图）
        max_images = int(cfg.get("media_max_images", 1))
        if images and max_images > 0 and self.services.openai_ready:
            for img in images[:max_images]:
                try:
                    desc = await self.services.describe_image(img, is_url=True)
                    notes.append(f"[页面内图片 {img}]\n{desc}")
                except svc.ServiceUnavailable as exc:
                    logger.warning("页面图片识别跳过 %s: %s", img, exc)

        # 5) 页面里挂着的文档（PDF/Office 附件、内嵌预览）
        max_docs = int(cfg.get("media_max_docs", 1))
        if docs and max_docs > 0 and self.services.document_model:
            for doc in docs[:max_docs]:
                try:
                    desc = await self._read_document_url(doc)
                    notes.append(f"[页面内文档 {doc}]\n{desc}")
                except svc.ServiceUnavailable as exc:
                    logger.warning("页面文档未能分析 %s: %s", doc, exc)
                    notes.append(f"[页面内文档 {doc} 未能分析：{exc}]")

        return "\n".join(notes) if notes else None

    # ------------------------- 输入归一化 -------------------------
    async def _normalize(self, user_input: dict, call_api, entry: dict | None = None) -> str:
        """把一轮输入归一化成纯文本。

        entry 非空时，每份媒体会被登记成一个短编号（m1/m2…），文本里以
        `[图片 m1 · 内容]` 的形式出现——模型据此在后续追问中引用原资源。
        """
        parts = []
        msg_id = user_input.get("msg_id")
        seg_map = user_input.get("segs") or {}

        def _loc(key: str, idx: int):
            """取该媒体在**本条消息**里的同类段序号；引用/转发来的返回 (None, None)。"""
            if entry is None or not msg_id:
                return None, None
            seq = seg_map.get(key) or []
            if idx >= len(seq):
                return None, None
            s = seq[idx]
            if isinstance(s, int) and s >= 0:
                return msg_id, s
            return None, None

        text = (user_input.get("text") or "").strip()
        if text:
            # 自动抓取文本里的链接，让 bot 能真正"打开"链接看内容
            urls = re.findall(r'https?://[^\s，。、；）)\'"<>]+', text)
            for u in urls[:3]:
                try:
                    # 媒体直链（视频/图片/文档/音频）→ 交给对应多模态模型"看/听"；
                    # 非媒体直链（None）→ 抓网页文本
                    media = await self._describe_media_url(u)
                    if media is not None:
                        parts.append(media)
                        continue
                    content = await self.services.fetch_url(u, web_cfg=self.services.web_config)
                    parts.append(f"[链接内容 {u}]\n{content}")
                    # 页面里若嵌着视频/音频/图片 → 一并交给模型看/听
                    page_media = await self._describe_page_media(u)
                    if page_media:
                        parts.append(page_media)
                except Exception as exc:  # noqa: BLE001
                    logger.warning("抓取链接失败 %s: %s", u, exc)
            parts.append(text)

        ready = self.services.openai_ready
        # 首轮就把用户真正想问的带给多模态模型，避免"先产一段通用描述、
        # 再靠 inspect_media 补一次"的两次往返。没有可用问题时为 None（用默认 prompt）。
        ask = self._first_prompt(self._clean_question(user_input.get("text") or ""))

        # 图片段（image 消息段）→ 视觉理解
        for i, img in enumerate(user_input.get("images", [])):
            mid, seg = _loc("images", i)
            sink: dict = {}
            if ready:
                try:
                    desc = await self.services.describe_image(img, prompt=ask, sink=sink)
                    h = self._register_media(entry, "image", sink.get("ref"), msg_id=mid, seg=seg)
                    parts.append(f"[{self._label('image', h)}] {desc}")
                    continue
                except Exception as exc:  # noqa: BLE001
                    # 这里刻意宽于 ServiceUnavailable：上传/下载/解码抛的是原生
                    # URLError、OSError、socket.timeout，窄捕获会让它们穿透到
                    # main，用户就只剩一句"处理出错了"。按"这份媒体没看成"降级即可。
                    logger.warning("视觉理解跳过: %s: %s", type(exc).__name__, exc)
                    # 即使理解失败，也登记引用：之后用户追问时还能再试一次
                    h = self._register_media(entry, "image", sink.get("ref") or {"file_id": ""},
                                             msg_id=mid, seg=seg)
                    parts.append(f"[{self._label('image', h, '识别失败')}] 视觉理解失败（{_reason(exc)}）")
                    continue
            parts.append("[用户发送了图片，但当前未配置视觉服务，无法识别内容]")

        # 语音段（record 消息段）→ 优先方舟 Chat 音频理解（input_audio，走 audio_model）；失败回退 ASR 转写
        for i, voc in enumerate(user_input.get("voices", [])):
            mid, seg = _loc("voices", i)
            sink = {}
            if self.services.audio_model:
                try:
                    desc = await self.services.understand_audio(voc, prompt=ask, sink=sink)
                    h = self._register_media(entry, "audio", sink.get("ref"), msg_id=mid, seg=seg)
                    parts.append(f"[{self._label('audio', h)}] {desc}")
                    continue
                except Exception as exc:  # noqa: BLE001
                    logger.warning("音频模型理解失败，回退 ASR 转写: %s: %s",
                                   type(exc).__name__, exc)
            h = self._register_media(entry, "audio", sink.get("ref") or {"file_id": ""},
                                     msg_id=mid, seg=seg)
            if self.services.voice_ready:
                try:
                    txt = await self.services.speech_to_text(voc)
                    parts.append(f"[{self._label('audio', h, '转写')}] {txt}")
                    continue
                except Exception as exc:  # noqa: BLE001
                    logger.warning("语音识别跳过: %s: %s", type(exc).__name__, exc)
            parts.append(f"[{self._label('audio', h)}] 语音识别不可用，无法分析内容")

        # 视频段（video 消息段）→ 豆包原生视频理解
        for i, vid in enumerate(user_input.get("videos", [])):
            mid, seg = _loc("videos", i)
            sink = {}
            if self.services.video_ready:
                try:
                    desc = await self.services.understand_video(vid, prompt=ask, sink=sink)
                    h = self._register_media(entry, "video", sink.get("ref"),
                                             msg_id=mid, seg=seg,
                                             audio_ref=sink.get("audio_ref"))
                    parts.append(f"[{self._label('video', h)}] {desc}")
                    continue
                except Exception as exc:  # noqa: BLE001
                    logger.warning("视频理解失败: %s: %s", type(exc).__name__, exc)
                    h = self._register_media(entry, "video", sink.get("ref") or {"file_id": ""},
                                             msg_id=mid, seg=seg)
                    parts.append(f"[{self._label('video', h, '识别失败')}] 视频理解失败：{_reason(exc)}")
                    continue
            parts.append("[用户发送了视频，但当前未配置视频理解服务，无法分析内容]")

        # 文件（上传：group_upload / offline_file）→ 按扩展名路由到对应"手"
        for i, (name, path) in enumerate(user_input.get("files", [])):
            mid, seg = _loc("files", i)
            try:
                parts.append(await self._handle_file(name, path, ready, prompt=ask,
                                                     entry=entry, msg_id=mid, seg=seg))
            except Exception as exc:  # noqa: BLE001
                logger.warning("文件处理失败 %s: %s: %s", name, type(exc).__name__, exc)
                parts.append(f"[文件 {name} 未能分析：{_reason(exc)}]")

        return "\n".join(p for p in parts if p)

    @staticmethod
    def _clean_question(text: str) -> str:
        """从本轮文本里提炼"用户想问什么"，用于拼进首轮多模态 prompt。

        要剥掉三类噪音，否则会污染视觉模型的输入：
          1. `[引用的消息 · …]`  ——是被引用者的原话，不是本轮的问题
          2. 群聊的 `发送者：` 前缀（main 里拼上的，且该媒体不是他发的）
          3. `[分享链接]` 行与 URL 本身
        返回空串表示"没有可用问题"，调用方会退回各类型的默认 prompt。
        """
        if not text:
            return ""
        lines = text.splitlines()

        # 1) 先剥群聊发送者前缀（main 里拼在**首行**：`御坂：原文本`）。
        #    必须先剥，否则下面的 [引用的消息] 过滤会因为行首多了名字而失效。
        if lines:
            m = re.match(r"^([^：:\n]{1,12})[：:]\s*(.*)$", lines[0])
            left = m.group(1) if m else ""
            # 昵称里不该出现这些字符；否则 "[引用的消息 · 张三：…]" 会被误当成"昵称：内容"
            if (m and "/" not in left and not left.lower().startswith("http")
                    and not any(ch in left for ch in "[]（）()·、，,。")):
                lines[0] = m.group(2)

        # 2) 合并转发的聊天记录整体丢弃：它是一大段别人说的话，不是本轮的问题，
        #    塞进首轮 prompt 会让视觉模型对着几十行聊天记录去"回答"。
        cut = next((i for i, ln in enumerate(lines)
                    if ln.strip().startswith("[合并转发")), None)
        if cut is not None:
            lines = lines[:cut]

        # 3) 逐行过滤
        keep = []
        for ln in lines:
            s = ln.strip()
            if not s or s.startswith("[引用的消息") or s.startswith("[分享链接]"):
                continue
            keep.append(s)
        q = " ".join(keep).strip()
        q = re.sub(r"https?://\S+", "", q).strip()
        if len(q) < 2:
            return ""
        # 短句若不像提问（如"哈哈哈哈"）就别硬塞给视觉模型。
        # 中文三字问句很常见（"怎么样""谁来了"），所以下限设到 2 而不是 4。
        hints = ("什么", "哪", "谁", "几", "怎么", "为什么", "吗", "?", "？",
                 "说说", "看看", "描述", "总结", "翻译", "讲", "认", "读")
        if len(q) < 8 and not any(w in q for w in hints):
            return ""
        return q

    @staticmethod
    def _first_prompt(question: str) -> str | None:
        """首轮多模态 prompt：固定约束 + 用户的实际问题。

        "先答问题、再补概述"是刻意的——只答问题会丢掉画面整体信息，
        之后用户问别的就得再调一次 inspect_media；只给概述则等于没改。
        没有可用问题时返回 None，让 services 用各自调优过的默认 prompt。
        """
        if not question:
            return None
        return (
            f"用户的问题是：{question}\n"
            "请先直接回答这个问题；再补充一两句该媒体的整体内容"
            "（主要人物、场景、可见文字或声音）。"
            "如果问题涉及的信息在这个媒体里确实看不到/听不到，"
            "请明确说「看不到」，不要凭常理猜测。")

    @staticmethod
    def _label(kind: str, handle: str | None, suffix: str = "内容") -> str:
        """对话里的媒体标记：有编号写「图片 m1 · 内容」，无编号退回旧的「图片内容」。"""
        label = _MEDIA_LABEL.get(kind, "媒体")
        return f"{label} {handle} · {suffix}" if handle else f"{label}{suffix}"

    async def _handle_file(self, name: str, path: str, ready: bool, *,
                           prompt: str | None = None,
                           entry: dict | None = None, msg_id=None, seg=None) -> str:
        """把上传的文件按类型交给对应服务"查看"。

        prompt 为 None 时各服务用自己调优过的默认提示；传入则由
        `_first_prompt` 生成（含用户本轮真正想问的问题）。
        """
        ext = Path(name).suffix.lower()
        note = f"[用户发送了文件：{name}]"
        # 0) 视频 → 豆包原生视频理解（容器型扩展名）
        if ext in _VIDEO_EXTS:
            if self.services.video_ready:
                sink: dict = {}
                try:
                    desc = await self.services.understand_video(path, prompt=prompt, sink=sink)
                    h = self._register_media(entry, "video", sink.get("ref"),
                                             msg_id=msg_id, seg=seg,
                                             audio_ref=sink.get("audio_ref"))
                    return (f"{note}\n[{self._label('video', h)}] {desc}" if desc
                            else f"{note}\n（视频中未识别到内容）")
                except svc.ServiceUnavailable as exc:
                    self._register_media(entry, "video", sink.get("ref") or {"file_id": ""},
                                         msg_id=msg_id, seg=seg)
                    return f"{note}\n（视频理解不可用：{exc}）"
            return f"{note}\n（未配置视频理解服务，无法分析视频）"
        # 1) 音频 → 优先方舟 Chat 音频理解（input_audio，走 audio_model）；失败回退 ASR 转写
        if ext in _AUDIO_EXTS:
            sink = {}
            if self.services.audio_model:
                try:
                    desc = await self.services.understand_audio(path, prompt=prompt, sink=sink)
                    h = self._register_media(entry, "audio", sink.get("ref"),
                                             msg_id=msg_id, seg=seg)
                    return (f"{note}\n[{self._label('audio', h)}] {desc}" if desc
                            else f"{note}\n（音频中未识别到内容）")
                except svc.ServiceUnavailable as exc:
                    logger.warning("音频模型理解失败，回退 ASR: %s", exc)
            h = self._register_media(entry, "audio", sink.get("ref") or {"file_id": ""},
                                     msg_id=msg_id, seg=seg)
            if self.services.voice_ready:
                try:
                    txt = await self.services.speech_to_text(path)
                    return (f"{note}\n[{self._label('audio', h, '转写')}] {txt}" if txt
                            else f"{note}\n（音频中未识别到文字）")
                except svc.ServiceUnavailable as exc:
                    return f"{note}\n（语音识别不可用：{exc}）"
            return f"{note}\n（未配置音频理解/语音识别服务，无法处理音频）"
        # 2) 图片 → 视觉理解
        if ext in _IMAGE_EXTS:
            if ready:
                sink = {}
                try:
                    desc = await self.services.describe_image(path, prompt=prompt, sink=sink)
                    h = self._register_media(entry, "image", sink.get("ref"),
                                             msg_id=msg_id, seg=seg)
                    return f"{note}\n[{self._label('image', h)}] {desc}"
                except svc.ServiceUnavailable as exc:
                    # 上传成功但理解失败时，引用仍然有效，登记后以后还能追问
                    self._register_media(entry, "image", sink.get("ref") or {"file_id": ""},
                                         msg_id=msg_id, seg=seg)
                    return f"{note}\n（视觉服务不可用：{exc}）"
            return f"{note}\n（未配置视觉服务，无法识别图片）"
        # 3) 文档（PDF/Office）→ 方舟 Chat 的 file 内容类型做多模态文档理解（走 document_model）；
        #    失败回退本地高精度解析
        if ext in _DOC_MODEL_EXTS:
            sink = {}
            if self.services.document_model:
                try:
                    desc = await self.services.understand_document(path, prompt=prompt, sink=sink)
                    h = self._register_media(entry, "document", sink.get("ref"),
                                             msg_id=msg_id, seg=seg)
                    return (f"{note}\n[{self._label('document', h)}] {desc}" if desc
                            else f"{note}\n（文档中未识别到内容）")
                except svc.ServiceUnavailable as exc:
                    logger.warning("文档模型理解失败，回退本地解析: %s", exc)
                    # 文档已上传成功时引用仍有效，登记后仍可用 inspect_media 追问
                    self._register_media(entry, "document", sink.get("ref") or {"file_id": ""},
                                         msg_id=msg_id, seg=seg)
            content = _read_file_text(path)
            if content is not None:
                return f"{note}\n文件内容（本地解析）：\n{content}"
            return f"{note}\n（未配置文档理解模型，且本地解析失败）"
        # 4) 文本类 → 直接本地读取（txt/md/csv 等 Files API 不支持，也不上传）
        content = _read_file_text(path)
        if content is not None:
            return f"{note}\n文件内容：\n{content}"
        return f"{note}\n（该文件类型暂不支持解析，已收到文件）"

    # ------------------------- 消息组装 -------------------------
    def _build_messages(self, entry: dict, user_text: str) -> list:
        messages = [{"role": "system", "content": self._system_prompt()}]
        summary = entry.get("summary", "") or ""
        if summary:
            messages.append(
                {"role": "system", "content": "【与该用户的长期记忆摘要】\n" + summary}
            )
        messages.extend(entry.get("recent", []))
        messages.append({"role": "user", "content": user_text})
        return messages

    def _system_prompt(self) -> str:
        base = self.llm.system_prompt
        extra = (
            "\n你可以使用工具：web_search（查询最新/实时信息）；"
            "fetch_url（打开并读取用户发来的链接/网页内容）；"
            "inspect_media（重新查看本会话里发过的图片/视频/音频/文档，追问细节）；"
            "generate_image（用户要你画图或发图时生成图片）；"
            "text_to_speech（用户要你用语音回复时合成语音）。"
            "用户发来链接时，应主动用 fetch_url 打开阅读后再回答；按需调用，不要虚构工具结果。"
            "\n历史消息里的 [图片 m1] / [视频 m2] / [音频 m3] / [文档 m4] 是内部媒体编号，"
            "用户追问某个媒体的细节而你手上的描述不够时，用 inspect_media 带上编号和具体问题去重看；"
            "这些编号和 [mm-dd hh:mm] 时间戳都是内部标记，不要向用户复述或解释。"
        )
        return base + extra

    # ------------------------- 工具循环 -------------------------
    async def _tool_loop(self, messages: list, tools: list, outputs: dict,
                         entry: dict | None = None, call_api=None) -> str:
        for round_no in range(MAX_TOOL_ROUNDS):
            try:
                resp = await self.llm.complete(messages, tools)
            except Exception as exc:  # noqa: BLE001
                # 带堆栈记录原始异常；给用户只回一句人话，不泄露 key/端点等内部细节
                logger.exception("LLM 调用失败: %s", exc)
                return self._llm_fail_note(exc)
            choices = resp.get("choices") or []
            if not choices:
                # 内容审核拦截、上游限流等都可能返回空 choices；
                # 不接住的话下一行就是 IndexError，用户只能收到一句"处理出错"
                logger.warning("LLM 返回空 choices（可能被内容审核拦截）: %s", str(resp)[:300])
                return "（我这轮没能生成回复：模型没有返回内容，可能触发了内容审核，换个说法再试。）"
            choice_obj = choices[0] or {}
            choice = choice_obj.get("message") or {}
            finish = choice_obj.get("finish_reason") or ""
            tool_calls = choice.get("tool_calls") or []
            if tool_calls:
                # 必须在下一次请求之前补齐：漏 id 会让整轮被判消息结构不合法
                _ensure_tool_call_ids(tool_calls, round_no)
            messages.append(choice)  # 含可能的 tool_calls，供下一轮上下文

            if not tool_calls:
                return (choice.get("content") or "").strip()

            for tc in tool_calls:
                fn = tc.get("function") or {}
                name = fn.get("name") or ""
                raw_args = fn.get("arguments")
                args, why = _parse_tool_args(raw_args, name, finish)
                if why:
                    # 不执行，但**必须**回一条 tool 消息：每个 tool_call 都要有对应的
                    # tool 回复，否则下一轮请求本身就会被上游判为消息结构不合法。
                    snippet = (" ".join(str(raw_args).split())[:120]
                               if raw_args else "（空）")
                    logger.warning(
                        "工具参数无效，本轮不执行: tool=%s finish=%s 原因=%s 参数=%s",
                        name or "?", finish or "-", why, snippet)
                    result_text = (f"（工具 {name or '未知'} 本次未执行：{why}。"
                                   f"你发来的参数是：{snippet}。"
                                   "请修正后重新调用，或直接用文字回答用户。）")
                else:
                    result_text = await self._exec_tool(name, args, outputs, entry, call_api)
                messages.append({
                    "role": "tool",
                    "tool_call_id": tc.get("id"),
                    "content": result_text,
                })
        return "（这轮工具调用次数用完了还没得出结论，可以换个更具体的问法再试一次。）"

    @staticmethod
    def _llm_fail_note(exc: BaseException) -> str:
        """把 LLM 异常翻成一句用户能看懂的话（不含内部细节）。

        之所以要分类：超时/限流是"稍后重试就能好"，鉴权失败是"重试也没用"，
        笼统回一句"出错了"会让人一直重试。
        """
        detail = f"{type(exc).__name__}: {exc}".lower()
        if isinstance(exc, (asyncio.TimeoutError, TimeoutError)) or "timed out" in detail:
            why = "模型响应超时"
        elif "429" in detail or "rate limit" in detail or "too many" in detail:
            why = "模型服务限流"
        elif "401" in detail or "403" in detail or "unauthorized" in detail:
            why = "模型鉴权失败"
        else:
            why = "模型服务异常"
        return f"（我这轮没能生成回复：{why}，可以稍后再试一次。）"

    async def _exec_tool(self, name: str, args: dict, outputs: dict,
                         entry: dict | None = None, call_api=None) -> str:
        try:
            if name == "web_search":
                snippets = await self.services.web_search(args.get("query", ""))
                return f"搜索结果：\n{snippets}"
            if name == "fetch_url":
                url = args.get("url", "")
                # 媒体直链：转给对应多模态模型（否则按网页抓文本只会拿到二进制乱码）
                media = await self._describe_media_url(url)
                if media is not None:
                    return f"该链接是媒体直链，已用多模态模型解析：\n{media}"
                content = await self.services.fetch_url(
                    url, web_cfg=self.services.web_config)
                # 网页里内嵌的视频/音频/图片也一并分析
                page_media = await self._describe_page_media(url)
                if page_media:
                    return f"网页内容：\n{content}\n\n{page_media}"
                return f"网页内容：\n{content}"
            if name == "inspect_media":
                return await self._inspect_media(
                    entry, args.get("handle", ""), args.get("question", ""), call_api)
            if name == "generate_image":
                path = await self.services.generate_image(args.get("prompt", ""))
                outputs["images"].append(path)
                return f"已生成图片（本地路径：{path}），将以图片消息发出。"
            if name == "text_to_speech":
                path = await self.services.text_to_speech(args.get("text", ""))
                outputs["voices"].append(path)
                return f"已合成语音（本地路径：{path}），将以语音消息发出。"
        except svc.ServiceUnavailable as exc:
            return f"（工具 {name} 不可用：{exc}）"
        except Exception as exc:  # noqa: BLE001
            logger.error("工具 %s 执行失败: %s", name, exc)
            return f"（工具 {name} 执行失败：{exc}）"
        return f"（未知工具：{name}）"

    # ------------------------- 输出拼装 -------------------------
    def _build_segments(self, reply_text: str, outputs: dict) -> list:
        reply_text = reply_text[: self.max_reply_length]
        segments = [{"type": "text", "data": {"text": reply_text}}]
        for img in outputs["images"]:
            segments.append({"type": "image", "data": {"file": img}})
        for voc in outputs["voices"]:
            segments.append({"type": "record", "data": {"file": voc}})
        for name, path in outputs["files"]:
            segments.append({"type": "file", "data": {"name": name, "file": path}})
        return segments
